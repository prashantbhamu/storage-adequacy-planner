from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .data_io import ValidatedInput
from .optimizer import METHOD_OBJECTIVES, VERSION, OptimizationResult, StorageSpec, storage_spec_dict


NAVY = "0B2B50"
TEAL = "078C82"
PALE_TEAL = "E9F7F5"
PALE_BLUE = "EEF5FA"
SLATE = "5D6B7A"
LIGHT_BORDER = "C9D6E2"
WHITE = "FFFFFF"


def _section_title(sheet, row: int, text: str, last_column: int = 6) -> None:
    cell = sheet.cell(row, 1, text)
    cell.font = Font(name="Aptos Display", size=14, bold=True, color=NAVY)
    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    cell.alignment = Alignment(vertical="center")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)


def _label(cell, text: str) -> None:
    cell.value = text
    cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)


def _header_row(sheet, row: int, headers: tuple[str, ...]) -> None:
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)


def _key_values(sheet, row: int, rows: list[tuple], number_format: str = "0.000") -> int:
    """Write label/value/unit rows; return the next free row."""
    for label, value, *unit in rows:
        _label(sheet.cell(row, 1), label)
        cell = sheet.cell(row, 2, value)
        if isinstance(value, float):
            cell.number_format = unit[1] if len(unit) > 1 else number_format
        if unit:
            sheet.cell(row, 3, unit[0])
        row += 1
    return row + 1


def _style_summary_sheet(
    workbook: Workbook,
    validated: ValidatedInput,
    spec: StorageSpec,
    result: OptimizationResult,
) -> None:
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Storage Adequacy Planner"
    sheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:F2")
    sheet["A2"] = (
        f"{validated.period_label} · 48 h look-ahead · 24 h commitment · "
        "anchored to the perfect-foresight optimum"
    )
    sheet["A2"].font = Font(name="Aptos", size=10, color=SLATE)

    row = 4
    _section_title(sheet, row, "Run information")
    row = _key_values(sheet, row + 1, [
        ("Source file", validated.filename),
        ("Source worksheet", validated.sheet_name or "CSV"),
        ("Period", validated.period_label),
        ("Hours", len(validated.timestamps)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Tool version", VERSION),
        ("Dispatch policy", "Progressive leximin"),
    ])

    assumptions = storage_spec_dict(spec)
    _section_title(sheet, row, "Storage assumptions")
    row = _key_values(sheet, row + 1, [
        ("Maximum charge", assumptions["charge_power_gw"], "GW"),
        ("Maximum discharge", assumptions["discharge_power_gw"], "GW"),
        ("Energy capacity", assumptions["energy_gwh"], "GWh"),
        ("Round-trip efficiency", assumptions["rte"], "%", "0.000%"),
        ("One-way efficiency", assumptions["charge_efficiency"], "%", "0.000%"),
        ("Maximum cycles / 06:00 day", assumptions["max_cycles_per_accounting_day"], "cycles"),
        ("Internal throughput cap", assumptions["daily_internal_throughput_cap_gwh"], "GWh/direction"),
        ("Initial SOC", assumptions["initial_soc_fraction"], "% of capacity", "0.0%"),
        ("Final SOC", assumptions["final_soc_fraction"], "% of capacity", "0.0%"),
        ("Minimum SOC", assumptions["min_soc_fraction"], "% of capacity", "0.0%"),
        ("Maximum SOC", assumptions["max_soc_fraction"], "% of capacity", "0.0%"),
        ("Charging allowed", "Surplus hours only" if spec.charge_from_surplus_only else "Any hour"),
    ])

    summary = result.summary
    _section_title(sheet, row, "Before and after")
    _header_row(sheet, row + 1, ("Metric", "Before", "After", "Change", "Unit", "Interpretation"))
    row += 2
    for label, before, after, unit, interpretation in (
        ("Minimum residual gap", summary["minimum_gap_before_gw"], summary["minimum_gap_after_gw"], "GW", "Higher is better"),
        ("Maximum residual gap", summary["maximum_gap_before_gw"], summary["maximum_gap_after_gw"], "GW", "Lower after charging is better"),
        ("Shortage energy", summary["shortage_energy_before_gwh"], summary["shortage_energy_after_gwh"], "GWh", "Lower is better"),
        ("Shortage hours", summary["shortage_hours_before"], summary["shortage_hours_after"], "hours", "Lower is better"),
    ):
        for column, value in enumerate((label, before, after, after - before, unit, interpretation), start=1):
            cell = sheet.cell(row, column, value)
            if column in (2, 3, 4):
                cell.number_format = "0.000"
        row += 1
    row += 1

    benchmark = result.benchmark
    _section_title(sheet, row, "Perfect-foresight check")
    row = _key_values(sheet, row + 1, [
        ("Perfect-foresight floor", benchmark["perfect_foresight_floor_gw"], "GW"),
        ("Rolling-horizon floor", benchmark["rolling_floor_gw"], "GW"),
        ("Perfect-foresight shortage", benchmark["perfect_foresight_shortage_gwh"], "GWh"),
        ("Rolling-horizon shortage", benchmark["rolling_shortage_gwh"], "GWh"),
        ("Floor below optimum", benchmark["floor_shortfall_gw"], "GW", "0.000000"),
        ("Shortage above optimum", benchmark["excess_shortage_gwh"], "GWh", "0.000000"),
    ])

    limits = result.limits
    _section_title(sheet, row, "What limits the result")
    _header_row(sheet, row + 1, ("Binding limit", "Shortage hours", "Shortage energy (GWh)"))
    row += 2
    for key, label in limits["labels"].items():
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, limits["shortage_hours"][key])
        sheet.cell(row, 3, limits["shortage_energy_gwh"][key]).number_format = "0.000"
        row += 1
    floor_hour = limits["floor_hour"]
    _label(sheet.cell(row, 1), "Lowest residual hour")
    sheet.cell(row, 2, floor_hour["timestamp"])
    sheet.cell(row, 3, floor_hour["residual_gap_gw"]).number_format = "0.000"
    sheet.cell(row, 4, limits["labels"][floor_hour["limit"]])
    row += 1
    _label(sheet.cell(row, 1), "Most effective +10% increase")
    sheet.cell(row, 2, limits["most_effective_increase"] or "None of the tested limits")
    row += 2

    if result.sensitivity:
        _section_title(sheet, row, "Sensitivity: +10% of each limit (perfect foresight)")
        _header_row(sheet, row + 1, ("Parameter", "From", "To", "Unit", "Floor change (GW)", "Shortage change (GWh)"))
        row += 2
        for item in result.sensitivity:
            values = (item["parameter"], item["from"], item["to"], item["unit"],
                      item["floor_change_gw"], item["shortage_change_gwh"])
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column, value)
                if isinstance(value, float):
                    cell.number_format = "0.000"
            row += 1
        row += 1

    _section_title(sheet, row, "Storage operation")
    row = _key_values(sheet, row + 1, [
        ("Peak charging", summary["peak_charge_gw"], "GW"),
        ("Peak discharging", summary["peak_discharge_gw"], "GW"),
        ("Total grid-side charging", summary["total_charge_gwh"], "GWh"),
        ("Total grid-side discharging", summary["total_discharge_gwh"], "GWh"),
        ("Conversion losses", summary["conversion_losses_gwh"], "GWh"),
        ("Equivalent full cycles", summary["equivalent_cycles"], "cycles"),
    ])

    diagnostics = result.validation["diagnostics"]
    checks = result.validation["checks"]
    _section_title(sheet, row, "Constraint validation")
    row = _key_values(sheet, row + 1, [
        ("Overall validation", "Passed" if result.validation["passed"] else "Failed"),
        ("Minimum SOC", diagnostics["minimum_soc_gwh"], "GWh", "0.000000"),
        ("Maximum SOC", diagnostics["maximum_soc_gwh"], "GWh", "0.000000"),
        ("Final SOC", diagnostics["final_soc_gwh"], "GWh", "0.000000"),
        ("Maximum charge", diagnostics["maximum_charge_gw"], "GW", "0.000000"),
        ("Maximum discharge", diagnostics["maximum_discharge_gw"], "GW", "0.000000"),
        ("Maximum internal charge / 06:00 day", diagnostics["maximum_06_day_internal_charge_gwh"], "GWh", "0.000000"),
        ("Maximum internal discharge / 06:00 day", diagnostics["maximum_06_day_internal_discharge_gwh"], "GWh", "0.000000"),
        ("Simultaneous charge/discharge", checks["simultaneous_charge_discharge_gw"], "GW", "0.000000"),
        ("Charging beyond surplus", checks["surplus_charging_violation_gw"], "GW", "0.000000"),
    ])

    _section_title(sheet, row, "Daily performance")
    _header_row(sheet, row + 1, (
        "Accounting day",
        "Minimum gap before (GW)",
        "Minimum gap after (GW)",
        "Shortage before (GWh)",
        "Shortage after (GWh)",
        "Equivalent cycles",
    ))
    row += 2
    for daily in result.daily_performance:
        values = (
            daily["date"],
            daily["minimum_gap_before_gw"],
            daily["minimum_gap_after_gw"],
            daily["shortage_energy_before_gwh"],
            daily["shortage_energy_after_gwh"],
            daily["equivalent_cycles"],
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
            if column > 1:
                sheet.cell(row, column).number_format = "0.000"
        row += 1

    widths = {1: 38, 2: 20, 3: 22, 4: 30, 5: 20, 6: 30}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


HOURLY_HEADERS = (
    "Timestamp",
    "Demand (GW)",
    "Available Supply (GW)",
    "Raw Gap (GW)",
    "Storage Charge (GW)",
    "Storage Discharge (GW)",
    "Storage Dispatch (GW)",
    "Adjusted Supply (GW)",
    "Residual Gap (GW)",
    "SOC End (GWh)",
    "06:00 Accounting Day",
)


def _style_hourly_sheet(workbook: Workbook, result: OptimizationResult) -> None:
    sheet = workbook.create_sheet("Hourly Results")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    for column, value in enumerate(HOURLY_HEADERS, start=1):
        cell = sheet.cell(1, column, value)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, timestamp in enumerate(result.timestamps, start=2):
        source = index - 2
        values = (
            timestamp,
            result.demand[source],
            result.supply[source],
            result.raw_gap[source],
            result.charge[source],
            result.discharge[source],
            result.dispatch[source],
            result.adjusted_supply[source],
            result.residual_gap[source],
            result.soc_end[source],
            result.accounting_days[source],
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(index, column, value)
        sheet.cell(index, 1).number_format = "yyyy-mm-dd hh:mm"
        for column in range(2, 11):
            sheet.cell(index, column).number_format = "0.000000"

    last_column = get_column_letter(len(HOURLY_HEADERS))
    table = Table(displayName="HourlyStorageResults", ref=f"A1:{last_column}{len(result.timestamps) + 1}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    widths = (20, 16, 22, 16, 20, 23, 23, 22, 20, 18, 22)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def build_results_workbook(
    validated: ValidatedInput,
    spec: StorageSpec,
    result: OptimizationResult,
) -> bytes:
    workbook = Workbook()
    workbook.properties.title = "Storage Adequacy Planner — Progressive Leximin"
    workbook.properties.description = (
        "; ".join(METHOD_OBJECTIVES)
        + ". Hour counts use a 1 kW numerical zero tolerance."
    )
    _style_summary_sheet(workbook, validated, spec, result)
    _style_hourly_sheet(workbook, result)
    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.outlinePr.summaryBelow = True
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row > 1 and cell.font == Font():
                    cell.font = Font(name="Aptos", size=10, color=NAVY)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
