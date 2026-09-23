import io
import json

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.data_io import detect_columns, validate_uploaded_table
from backend.main import app


def month_frame():
    timestamps = pd.date_range("2027-02-01", "2027-02-28 23:00", freq="h")
    hours = np.arange(len(timestamps)) % 24
    return pd.DataFrame({
        "Timestamp": timestamps,
        "Demand (GW)": np.full(len(timestamps), 10.0),
        "Available Supply (GW)": 10 + 5 * np.exp(-((hours - 13) / 3) ** 2) - 2.5 * np.exp(-((hours - 21) / 2) ** 2),
    })


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_three_column_upload_without_mapping(extension):
    frame = month_frame()
    if extension == "csv":
        payload = frame.to_csv(index=False).encode()
    else:
        stream = io.BytesIO()
        frame.to_excel(stream, index=False)
        payload = stream.getvalue()
    result = validate_uploaded_table(payload, f"input.{extension}")
    assert len(result.timestamps) == 672
    assert set(result.mapping) == {"timestamp", "demand", "supply"}


def test_aliases_case_and_whitespace():
    assert detect_columns(["Date Time", "LOAD (GW)", "Max Dispatch (GW)"]) == {
        "timestamp": "Date Time", "demand": "LOAD (GW)", "supply": "Max Dispatch (GW)",
    }


@pytest.mark.parametrize("headers,message", [
    (["Timestamp", "Demand (GW)"], "Missing recognised Available Supply"),
    (["Timestamp", "Demand (GW)", "Supply (GW)", "Available Supply (GW)"], "Ambiguous"),
    (["Timestamp", "Demand (MW)", "Available Supply (GW)"], "Power values must be in GW"),
    (["Timestamp", "High Demand (GW)", "Low Demand (GW)", "Supply (GW)"], "Missing recognised Demand"),
    (["Timestamp", "Demand (GW)", "Load (GW)", "Supply (GW)"], "Ambiguous"),
    (["Timestamp", "Demand (GW)", "Demand (GW).1", "Supply (GW)"], "Ambiguous"),
])
def test_missing_ambiguous_or_wrong_unit_headers_fail_closed(headers, message):
    with pytest.raises(ValueError, match=message):
        detect_columns(headers)


def test_extra_columns_such_as_solar_are_ignored():
    frame = month_frame()
    for extra in (1.5, "not required", -1):
        frame["Solar Generation (GW)"] = extra
        result = validate_uploaded_table(frame.to_csv(index=False).encode(), "input.csv")
        assert set(result.mapping) == {"timestamp", "demand", "supply"}


def test_bad_numeric_data_still_rejected():
    frame = month_frame()
    frame.loc[4, "Demand (GW)"] = np.nan
    with pytest.raises(ValueError, match="invalid numerical"):
        validate_uploaded_table(frame.to_csv(index=False).encode(), "input.csv")


SETTINGS = {"charge_power_gw": 2, "discharge_power_gw": 2, "energy_gwh": 8,
            "rte_percent": 90, "max_cycles_per_accounting_day": 1,
            "initial_soc_percent": 0, "final_soc_percent": 0}


def test_three_column_api_run_download_and_extra_column_parity():
    frame = month_frame()
    client = TestClient(app)
    files = {"file": ("february.csv", frame.to_csv(index=False).encode(), "text/csv")}
    validation = client.post("/api/validate", files=files)
    assert validation.status_code == 200
    settings = json.dumps(SETTINGS)
    response = client.post("/api/optimize", files=files, data={"settings": settings})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["validation"]["passed"]
    assert all("solar_gw" not in row for row in result["hourly"])
    download = client.get(f"/api/download/{result['run_id']}")
    assert download.status_code == 200
    workbook = load_workbook(io.BytesIO(download.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Hourly Results"]
    assert workbook["Hourly Results"].max_row == 673
    assert workbook["Hourly Results"].max_column == 11
    headers = next(workbook["Hourly Results"].iter_rows(max_row=1, values_only=True))
    assert not any("Solar" in str(value) for value in headers)
    workbook.close()
    frame["Solar Generation (GW)"] = 100.0
    extra = client.post("/api/optimize", files={"file": ("old.csv", frame.to_csv(index=False).encode(), "text/csv")}, data={"settings": settings})
    assert extra.status_code == 200
    for before, after in zip(result["hourly"], extra.json()["hourly"]):
        for key in ("charge_gw", "discharge_gw", "soc_end_gwh", "residual_gap_gw"):
            assert before[key] == after[key]


def test_api_requires_initial_and_final_soc():
    frame = month_frame()
    settings = {key: value for key, value in SETTINGS.items() if key != "final_soc_percent"}
    response = TestClient(app).post(
        "/api/optimize",
        files={"file": ("february.csv", frame.to_csv(index=False).encode(), "text/csv")},
        data={"settings": json.dumps(settings)},
    )
    assert response.status_code == 400
    assert "final_soc_percent" in response.json()["detail"]


def test_api_rejects_ambiguous_columns():
    frame = month_frame()
    frame["Supply (GW)"] = frame["Available Supply (GW)"]
    response = TestClient(app).post("/api/validate", files={"file": ("ambiguous.csv", frame.to_csv(index=False).encode(), "text/csv")})
    assert response.status_code == 400
    assert "Ambiguous" in response.json()["detail"]
