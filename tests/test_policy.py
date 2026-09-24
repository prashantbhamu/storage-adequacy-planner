"""User-defined SOC, surplus-only charging, prorated days, benchmark and sizing."""
from dataclasses import replace
from datetime import datetime, timedelta
import json
from pathlib import Path

import numpy as np
import pytest
from fastapi.testclient import TestClient

from backend.data_io import validate_uploaded_table
from backend.main import app
from backend.optimizer import (
    InfeasibleError,
    StorageSpec,
    _solve_horizon,
    accounting_day,
    accounting_day_caps,
    optimize_storage,
    perfect_foresight,
    suggest_cyclic_soc,
)
from backend.sizing import size_storage
from tests.test_tool import synthetic_run

APP_ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = APP_ROOT / "examples/synthetic_april_2031.csv"


@pytest.fixture(scope="module")
def april():
    return validate_uploaded_table(EXAMPLE.read_bytes(), EXAMPLE.name)


def test_user_defined_soc_levels_and_operating_range():
    times, demand, supply, base = synthetic_run()
    spec = replace(base, initial_soc_fraction=.25, final_soc_fraction=.75,
                   min_soc_fraction=.1, max_soc_fraction=.9)
    result = optimize_storage(times, demand, supply, spec, sensitivity=False)
    assert result.validation["passed"]
    assert abs(result.soc_end[-1] - 6.0) < 2e-6
    assert result.soc_end.min() >= 0.8 - 2e-6
    assert result.soc_end.max() <= 7.2 + 2e-6
    # The first hour starts from the user level, not from zero.
    first = 2.0 + spec.eta * result.charge[0] - result.discharge[0] / spec.eta
    assert abs(result.soc_end[0] - first) < 2e-6


@pytest.mark.parametrize("value,message", [
    (dict(initial_soc_fraction=.05, min_soc_fraction=.1), "Initial SOC"),
    (dict(final_soc_fraction=.95, max_soc_fraction=.9), "Final SOC"),
    (dict(min_soc_fraction=.6, max_soc_fraction=.5), "operating range"),
])
def test_soc_settings_are_validated(value, message):
    with pytest.raises(ValueError, match=message):
        replace(StorageSpec(2, 2, 8, .9, 1), **value).validate()


def test_surplus_only_charging_never_deepens_a_deficit():
    timestamps = [datetime(2031, 5, 1, 6) + timedelta(hours=i) for i in range(4)]
    gap = np.array([-1.0, -10.0, 5.0, 5.0])
    base = StorageSpec(100, 100, 20, 1, 1)
    caps = {accounting_day(timestamps[0]): 20.0}
    surplus_only = _solve_horizon(gap, timestamps, 0, {}, {}, base, caps, 0.0, final_soc=0.0)
    assert surplus_only["charge"][:2].max() < 1e-9
    assert abs(surplus_only["floor"] + 10.0) < 1e-6
    grid = _solve_horizon(gap, timestamps, 0, {}, {}, replace(base, charge_from_surplus_only=False),
                          caps, 0.0, final_soc=0.0)
    assert abs(grid["floor"] + 5.5) < 1e-6
    assert grid["charge"][0] > 4.4


def test_surplus_only_holds_over_a_full_run(april):
    spec = StorageSpec(40, 40, 200, .85, 1, initial_soc_fraction=.5, final_soc_fraction=.5)
    result = optimize_storage(april.timestamps, april.demand, april.supply, spec, sensitivity=False)
    assert np.all(result.charge <= np.maximum(result.raw_gap, 0) + 2e-6)
    assert result.validation["checks"]["surplus_charging_violation_gw"] <= 2e-6


def test_partial_accounting_days_are_prorated(april):
    spec = StorageSpec(40, 40, 200, .85, 1)
    caps = accounting_day_caps(april.timestamps, spec)
    assert caps["2031-03-31"] == pytest.approx(200 * 6 / 24)
    assert caps["2031-04-01"] == pytest.approx(200)
    assert caps["2031-04-30"] == pytest.approx(200 * 18 / 24)
    result = optimize_storage(april.timestamps, april.demand, april.supply, spec, sensitivity=False)
    first_day = [i for i, day in enumerate(result.accounting_days) if day == "2031-03-31"]
    assert sum(result.discharge[first_day]) / spec.eta <= 50 + 2e-6


def test_cyclic_suggestion_is_self_consistent(april):
    spec = StorageSpec(40, 40, 200, .85, 1)
    suggestion = suggest_cyclic_soc(april.timestamps, april.demand, april.supply, spec)
    assert 0 <= suggestion["soc_percent"] <= 100
    fraction = suggestion["soc_percent"] / 100
    cyclic = replace(spec, initial_soc_fraction=fraction, final_soc_fraction=fraction)
    plan = perfect_foresight(april.timestamps, april.supply - april.demand, cyclic)
    assert plan["floor_gw"] == pytest.approx(suggestion["floor_gw"], abs=1e-6)
    empty = perfect_foresight(april.timestamps, april.supply - april.demand, spec)
    assert suggestion["floor_gw"] >= empty["floor_gw"] - 1e-9


def test_unreachable_final_soc_fails_with_a_clear_message():
    times, demand, supply, base = synthetic_run()
    spec = replace(base, final_soc_fraction=1.0, max_cycles_per_accounting_day=.1)
    with pytest.raises(InfeasibleError, match="final SOC"):
        optimize_storage(times, demand, supply, spec)


def test_limits_cover_every_shortage_hour(april):
    spec = StorageSpec(40, 40, 200, .85, 1, initial_soc_fraction=.5, final_soc_fraction=.5)
    result = optimize_storage(april.timestamps, april.demand, april.supply, spec)
    assert sum(result.limits["shortage_hours"].values()) == result.summary["shortage_hours_after"]
    total = sum(result.limits["shortage_energy_gwh"].values())
    assert total == pytest.approx(result.summary["shortage_energy_after_gwh"], abs=1e-6)
    assert {row["parameter"] for row in result.sensitivity} == {
        "Charge power", "Discharge power", "Energy capacity", "Daily cycle limit"}


@pytest.mark.parametrize("mode,duration", [("energy", None), ("power", None), ("duration", 4.0)])
def test_sizing_is_minimal_and_meets_the_target(april, mode, duration):
    base = StorageSpec(40, 40, 200, .85, 1, initial_soc_fraction=.5, final_soc_fraction=.5)
    target = -35.0
    sized = size_storage(april.timestamps, april.demand, april.supply, base, mode, target, duration)
    assert sized["check_floor_gw"] >= target - 1e-6
    gap = april.supply - april.demand
    if mode == "energy":
        smaller = replace(base, energy_gwh=sized["energy_gwh"] * .99)
    else:
        power = sized["charge_power_gw"] * .99
        energy = power * duration if mode == "duration" else base.energy_gwh
        smaller = replace(base, charge_power_gw=power, discharge_power_gw=power, energy_gwh=energy)
    assert perfect_foresight(april.timestamps, gap, smaller)["floor_gw"] < target - 1e-6
    if mode == "duration":
        assert sized["energy_gwh"] == pytest.approx(4 * sized["charge_power_gw"])


def test_sizing_reports_unreachable_targets(april):
    base = StorageSpec(40, 40, 200, .85, 1)
    with pytest.raises(ValueError, match="discharge in the tightest hour"):
        size_storage(april.timestamps, april.demand, april.supply, base, "energy", 0.0)
    with pytest.raises(ValueError, match="No storage size"):
        size_storage(april.timestamps, april.demand, april.supply, base, "power", 1000.0)


def test_suggest_and_size_endpoints():
    client = TestClient(app, base_url="http://127.0.0.1")
    files = {"file": (EXAMPLE.name, EXAMPLE.read_bytes(), "text/csv")}
    settings = {"charge_power_gw": 40, "discharge_power_gw": 40, "energy_gwh": 200,
                "rte_percent": 85, "max_cycles_per_accounting_day": 1}
    suggestion = client.post("/api/suggest-soc", files=files, data={"settings": json.dumps(settings)})
    assert suggestion.status_code == 200, suggestion.text
    assert 0 <= suggestion.json()["soc_percent"] <= 100

    sizing_settings = {key: value for key, value in settings.items() if key != "energy_gwh"}
    sizing_settings.update(initial_soc_percent=50, final_soc_percent=50)
    sized = client.post("/api/size", files=files, data={
        "settings": json.dumps(sizing_settings),
        "sizing": json.dumps({"mode": "energy", "target_floor_gw": -35}),
    })
    assert sized.status_code == 200, sized.text
    assert sized.json()["energy_gwh"] > 0
    assert sized.json()["check_floor_gw"] >= -35 - 1e-6


def test_background_job_reports_progress_and_result():
    import time

    client = TestClient(app, base_url="http://127.0.0.1")
    files = {"file": (EXAMPLE.name, EXAMPLE.read_bytes(), "text/csv")}
    settings = {"charge_power_gw": 40, "discharge_power_gw": 40, "energy_gwh": 200,
                "rte_percent": 85, "max_cycles_per_accounting_day": 1,
                "initial_soc_percent": 50, "final_soc_percent": 50}
    started = client.post("/api/jobs", files=files, data={"settings": json.dumps(settings)})
    assert started.status_code == 200, started.text
    job_id = started.json()["job_id"]
    stages = set()
    for _ in range(600):
        status = client.get(f"/api/jobs/{job_id}").json()
        stages.add(status["stage"])
        if status["state"] != "running":
            break
        time.sleep(0.1)
    assert status["state"] == "done", status["error"]
    assert status["result"]["validation"]["passed"]
    assert status["done"] == status["total"]
    assert stages & {"benchmark", "horizons", "sensitivity"}
    assert client.get(f"/api/download/{status['result']['run_id']}").status_code == 200


def test_background_job_rejects_bad_settings_immediately():
    files = {"file": (EXAMPLE.name, EXAMPLE.read_bytes(), "text/csv")}
    response = TestClient(app, base_url="http://127.0.0.1").post("/api/jobs", files=files, data={"settings": json.dumps({"energy_gwh": 1})})
    assert response.status_code == 400


def test_validation_reports_raw_gap_statistics(april):
    response = TestClient(app, base_url="http://127.0.0.1").post("/api/validate", files={"file": (EXAMPLE.name, EXAMPLE.read_bytes(), "text/csv")})
    raw = response.json()["raw"]
    gap = april.supply - april.demand
    assert raw["minimum_gap_gw"] == pytest.approx(gap.min())
    assert raw["shortage_hours"] == 378
    assert raw["shortage_energy_gwh"] == pytest.approx(np.maximum(-gap, 0).sum())


def test_example_year_download_and_daily_preview():
    client = TestClient(app, base_url="http://127.0.0.1")
    example = client.get("/api/example")
    assert example.status_code == 200
    assert example.content == (APP_ROOT / "examples/synthetic_fy2029_30.csv").read_bytes()
    validation = client.post("/api/validate", files={"file": ("example.csv", example.content, "text/csv")}).json()
    assert validation["period_label"] == "FY 2029-30"
    assert validation["raw"]["shortage_hours"] > 1000
    days = validation["daily_minimum_gap_gw"]
    assert days[0]["day"] == "2029-03-31" and len(days) == 366


def test_only_local_pages_can_use_the_service():
    client = TestClient(app, base_url="http://127.0.0.1")
    assert client.get("/api/health").status_code == 200
    # DNS rebinding: another site's name pointed at this computer.
    assert client.get("/api/health", headers={"host": "attacker.example"}).status_code == 403
    # Another website posting through the visitor's browser.
    files = {"file": (EXAMPLE.name, EXAMPLE.read_bytes(), "text/csv")}
    foreign = client.post("/api/validate", files=files, headers={"origin": "https://attacker.example"})
    assert foreign.status_code == 403
    same = client.post("/api/validate", files=files, headers={"origin": "http://127.0.0.1"})
    assert same.status_code == 200


def test_oversized_uploads_are_refused(monkeypatch):
    import backend.main as main

    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 1000)
    response = TestClient(app, base_url="http://127.0.0.1").post(
        "/api/validate", files={"file": (EXAMPLE.name, EXAMPLE.read_bytes(), "text/csv")})
    assert response.status_code == 413


def test_workbook_keeps_uploaded_names_as_text():
    from io import BytesIO

    import pandas as pd
    from openpyxl import load_workbook

    from backend.data_io import ValidatedInput
    from backend.exporter import build_results_workbook

    timestamps, demand, supply, spec = synthetic_run()
    result = optimize_storage(timestamps, demand, supply, spec)
    validated = ValidatedInput(
        filename='=HYPERLINK("https://attacker.example","open").csv', sheet_name="=1+1",
        period_type="month", period_label="Synthetic", timestamps=timestamps, demand=demand,
        supply=supply, raw_frame=pd.DataFrame(), mapping={},
    )
    workbook = load_workbook(BytesIO(build_results_workbook(validated, spec, result)))
    cells = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row
             if isinstance(cell.value, str) and ("HYPERLINK" in cell.value or cell.value == "=1+1")]
    assert len(cells) == 2 and all(cell.data_type == "s" for cell in cells)
