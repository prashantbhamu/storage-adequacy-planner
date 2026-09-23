"""Leximin policy and public golden regression. No user-facing presets."""
from dataclasses import replace
from datetime import datetime, timedelta
import hashlib
import io
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from backend.data_io import validate_uploaded_table
from backend.exporter import build_results_workbook
from backend.optimizer import StorageSpec, _solve_horizon, accounting_day, optimize_storage
from tests.test_tool import synthetic_run

APP_ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("gaps,expected", [
    ([100., 5., 10., 20.], [80., 17.5, 17.5, 20.]),
    ([5., 100., 10., 20., 40.], [5., 80., 25., 25., 40.]),
    ([100., 90., 0., 0.], [85., 85., 10., 10.]),
])
def test_progressive_leveling_and_blocked_hours(gaps, expected):
    timestamps = [datetime(2031, 5, 1, 6) + timedelta(hours=i) for i in range(len(gaps))]
    spec = StorageSpec(100, 100, 20, 1, 1)
    gap = np.asarray(gaps)
    caps = {accounting_day(timestamps[0]): spec.daily_internal_throughput_cap_gwh}
    solved = _solve_horizon(gap, timestamps, 0, {}, {}, spec, caps, 0.0, final_soc=0.0)
    np.testing.assert_allclose(gap + solved["discharge"] - solved["charge"], expected, atol=1e-6)
    assert len(solved["levels"]) > 1
    assert max(np.minimum(solved["charge"], solved["discharge"])) < 1e-7


@pytest.mark.parametrize("spec", [
    StorageSpec(1, 3, 5, .5, .5),
    StorageSpec(3, 1, 5, .81, .75),
    StorageSpec(2, 2, 8, 1, 1.5),
    StorageSpec(.1, .1, .2, .9, 1),
    StorageSpec(100, 100, 1000, .9, 1),
])
def test_custom_storage_limits(spec):
    times, demand, supply, _ = synthetic_run()
    result = optimize_storage(times, demand, supply, spec, sensitivity=False)
    assert result.validation["passed"]
    assert result.validation["checks"]["simultaneous_charge_discharge_gw"] < 1e-7
    assert max(x["equivalent_cycles"] for x in result.daily_performance) <= spec.max_cycles_per_accounting_day + 2e-6
    assert result.benchmark["floor_shortfall_gw"] == 0


@pytest.mark.parametrize("field", ["charge_power_gw", "discharge_power_gw", "energy_gwh", "rte",
                                   "max_cycles_per_accounting_day", "initial_soc_fraction"])
@pytest.mark.parametrize("value", [float("nan"), float("inf")])
def test_nonfinite_assumptions_rejected(field, value):
    with pytest.raises(ValueError, match="finite"):
        replace(StorageSpec(2, 2, 8, .9, 1), **{field: value}).validate()


def test_public_synthetic_month_full_hourly_regression():
    """Pin every dispatch array to a deterministic, non-sensitive example."""
    path = APP_ROOT / "examples/synthetic_april_2031.csv"
    source = path.read_bytes()
    validated = validate_uploaded_table(source, path.name)
    assert len(validated.timestamps) == 720
    assert validated.timestamps[0] == datetime(2031, 4, 1)
    assert validated.timestamps[-1] == datetime(2031, 4, 30, 23)

    spec = StorageSpec(40, 40, 200, .85, 1, initial_soc_fraction=.5, final_soc_fraction=.5)
    result = optimize_storage(validated.timestamps, validated.demand, validated.supply, spec)

    def rounded_hash(values):
        payload = np.round(np.asarray(values, dtype="<f8"), 6).tobytes()
        return hashlib.sha256(payload).hexdigest()

    assert rounded_hash(result.charge) == "b36ed31bb3e543b73ffbf3104d5be8bd6e65536d8a6fa19f75f619108b7ebfcf"
    assert rounded_hash(result.discharge) == "0b02f2fb7bf01477e0f623607b939a58130945384e292dd7086c3445f380febe"
    assert rounded_hash(result.soc_end) == "2f5c5be9c173b6f379371dc5e24b28c90fe91243230d3b31d070663924dca631"
    assert rounded_hash(result.residual_gap) == "4486dad42dcd2d2696478224db85954af2fe0e978ed995e1bf72648e0d28e7a7"

    assert result.validation["passed"]
    assert result.summary["shortage_hours_before"] == 378
    assert result.summary["shortage_hours_after"] == 378
    assert abs(result.summary["shortage_energy_after_gwh"] - 5747.727726721794) < 2e-6
    assert abs(result.summary["equivalent_cycles"] - 29.499999997153008) < 2e-6
    assert result.validation["checks"]["simultaneous_charge_discharge_gw"] == 0
    assert result.benchmark["floor_shortfall_gw"] == 0
    assert result.benchmark["excess_shortage_gwh"] == 0
    assert result.limits["most_effective_increase"] == "Energy capacity"

    # Verify the complete downloadable workbook.
    workbook = load_workbook(io.BytesIO(build_results_workbook(validated, spec, result)))
    assert workbook.sheetnames == ["Summary", "Hourly Results"]
    hourly = workbook["Hourly Results"]
    assert hourly.max_row == 721 and hourly.max_column == 11
    assert hourly.tables["HourlyStorageResults"].ref == "A1:K721"
    rows = list(hourly.iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == validated.timestamps
    assert [row[10] for row in rows] == result.accounting_days
    np.testing.assert_allclose([row[7] for row in rows], result.adjusted_supply, atol=1e-10)
    np.testing.assert_allclose([row[8] for row in rows], result.residual_gap, atol=1e-10)
    np.testing.assert_allclose([row[9] for row in rows], result.soc_end, atol=1e-10)
    assert not any(cell.data_type in ("e", "f") for sheet in workbook for row in sheet for cell in row)
    summary = {row[0]: row[1:] for row in workbook["Summary"].iter_rows(values_only=True) if row[0]}
    assert summary["Shortage hours"][:2] == (378, 378)
    assert summary["Initial SOC"][0] == .5
    assert summary["Charging allowed"][0] == "Surplus hours only"
    assert summary["Most effective +10% increase"][0] == "Energy capacity"
    assert hashlib.sha256(path.read_bytes()).hexdigest() == hashlib.sha256(source).hexdigest()
    workbook.close()
