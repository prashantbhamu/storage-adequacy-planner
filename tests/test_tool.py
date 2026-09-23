from __future__ import annotations

import io
import json
from datetime import datetime, timedelta

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from backend.data_io import ValidatedInput, detect_and_validate_period
from backend.exporter import build_results_workbook
from backend.main import app
from backend.optimizer import StorageSpec, optimize_storage


def synthetic_run(hours: int = 72):
    timestamps = [datetime(2026, 1, 1) + timedelta(hours=index) for index in range(hours)]
    hour = np.arange(hours) % 24
    demand = np.full(hours, 10.0)
    supply = demand + 6.0 * np.exp(-((hour - 13.0) / 3.0) ** 2) - 3.0 * np.exp(-((hour - 21.0) / 2.0) ** 2)
    spec = StorageSpec(2.0, 2.0, 8.0, 0.9, 1.0)
    return timestamps, demand, supply, spec


def test_complete_month_detection() -> None:
    timestamps = pd.date_range("2028-02-01", "2028-02-29 23:00", freq="h")
    assert detect_and_validate_period(timestamps) == ("month", "February 2028")


def test_financial_year_detection_with_leap_day() -> None:
    timestamps = pd.date_range("2031-04-01", "2032-03-31 23:00", freq="h")
    period_type, label = detect_and_validate_period(timestamps)
    assert period_type == "financial_year"
    assert label == "FY 2031-32"
    assert len(timestamps) == 8784
    assert sum((timestamps.month == 2) & (timestamps.day == 29)) == 24


def test_partial_period_is_rejected() -> None:
    timestamps = pd.date_range("2026-05-02", periods=48, freq="h")
    try:
        detect_and_validate_period(timestamps)
    except ValueError as error:
        assert "complete calendar month" in str(error)
    else:
        raise AssertionError("Partial periods must be rejected")


def test_optimizer_constraints_and_objective_improvement() -> None:
    timestamps, demand, supply, spec = synthetic_run()
    result = optimize_storage(timestamps, demand, supply, spec)
    assert result.validation["passed"]
    assert result.validation["checks"]["simultaneous_charge_discharge_gw"] <= 2e-7
    assert result.validation["checks"]["final_soc_abs_gwh"] <= 2e-6
    assert result.summary["minimum_gap_after_gw"] >= result.summary["minimum_gap_before_gw"] - 2e-6
    assert result.summary["shortage_energy_after_gwh"] <= result.summary["shortage_energy_before_gwh"] + 2e-6
    assert max(row["equivalent_cycles"] for row in result.daily_performance) <= 1.0 + 2e-6


def test_clean_xlsx_export_contains_two_sheets() -> None:
    timestamps, demand, supply, spec = synthetic_run()
    result = optimize_storage(timestamps, demand, supply, spec)
    validated = ValidatedInput(
        filename="synthetic.csv",
        sheet_name=None,
        period_type="month",
        period_label="Synthetic",
        timestamps=timestamps,
        demand=demand,
        supply=supply,
        raw_frame=pd.DataFrame(),
        mapping={},
    )
    payload = build_results_workbook(validated, spec, result)
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    assert workbook.sheetnames == ["Summary", "Hourly Results"]
    assert workbook["Hourly Results"].max_row == len(timestamps) + 1
    assert workbook["Hourly Results"]["G2"].value == result.dispatch[0]
    assert workbook["Summary"]["B10"].value == "3.0.0"
    assert workbook["Summary"]["B11"].value == "Progressive leximin"


def test_xlsx_api_workflow_and_download() -> None:
    timestamps = pd.date_range("2027-02-01", "2027-02-28 23:00", freq="h")
    hour = np.arange(len(timestamps)) % 24
    demand = np.full(len(timestamps), 10.0)
    supply = demand + 5.0 * np.exp(-((hour - 13.0) / 3.0) ** 2) - 2.5 * np.exp(-((hour - 21.0) / 2.0) ** 2)
    source = pd.DataFrame({"Timestamp": timestamps, "Demand GW": demand, "Supply GW": supply})
    upload = io.BytesIO()
    source.to_excel(upload, index=False)
    payload = upload.getvalue()
    mapping = json.dumps({"timestamp": "Timestamp", "demand": "Demand GW", "supply": "Supply GW"})
    settings = json.dumps(
        {
            "charge_power_gw": 2.0,
            "discharge_power_gw": 2.0,
            "energy_gwh": 8.0,
            "rte_percent": 90.0,
            "max_cycles_per_accounting_day": 1.0,
            "initial_soc_percent": 50.0,
            "final_soc_percent": 50.0,
        }
    )
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    client = TestClient(app)
    assert client.get("/api/health").json()["version"] == "3.0.0"

    preview = client.post("/api/preview", files={"file": ("february.xlsx", payload, xlsx)})
    assert preview.status_code == 200
    assert preview.json()["row_count"] == 672

    validation = client.post(
        "/api/validate",
        files={"file": ("february.xlsx", payload, xlsx)},
        data={"mapping": mapping},
    )
    assert validation.status_code == 200
    assert validation.json()["period_label"] == "February 2027"

    optimisation = client.post(
        "/api/optimize",
        files={"file": ("february.xlsx", payload, xlsx)},
        data={"mapping": mapping, "settings": settings},
    )
    assert optimisation.status_code == 200
    result = optimisation.json()
    assert result["validation"]["passed"]
    assert result["period"]["hours"] == 672
    assert result["method"]["algorithm"] == "progressive leximin"
    assert result["method"]["objective"][-1] == "Minimise unnecessary throughput"
    assert result["method"]["charging"] == "surplus hours only"
    assert result["benchmark"]["floor_shortfall_gw"] == 0
    assert result["storage"]["initial_soc_gwh"] == 4.0
    assert abs(result["hourly"][-1]["soc_end_gwh"] - 4.0) < 2e-6
    assert set(result["limits"]["shortage_hours"]) == set(result["limits"]["labels"])

    download = client.get(f"/api/download/{result['run_id']}")
    assert download.status_code == 200
    assert "storage_dispatch_optimiser_February_2027" in download.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(download.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Hourly Results"]
    assert workbook["Hourly Results"].max_row == 673
